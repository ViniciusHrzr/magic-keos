#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
audit-docx.py — Audita fidelidade entre Magic no Universo Kéos v.0.4.docx e data/regras/

Uso:
  python scripts/audit-docx.py
  python scripts/audit-docx.py --verbose

Retorna:
  0 — sem drift detectado
  1 — drift encontrado (lista discrepâncias)
  2 — DOCX não encontrado
"""

import sys
import os
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

DOCX_PATH = "magic keos/Magic no Universo Kéos v.0.4.docx"
VERBOSE = "--verbose" in sys.argv


def extract_docx_text(path):
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            if any(cells):
                tables.append(cells)
    return paragraphs, tables


def check_anchor(docx_text_all, anchor):
    return any(anchor in line for line in docx_text_all)


def main():
    if not os.path.exists(DOCX_PATH):
        print(f"ERRO: DOCX não encontrado em {DOCX_PATH}")
        sys.exit(2)

    paragraphs, table_rows = extract_docx_text(DOCX_PATH)
    all_text = paragraphs + [' | '.join(row) for row in table_rows]

    # Âncoras críticas por seção — strings que DEVEM aparecer no docx
    # Derivadas de data/regras/ — representam conteúdo de jogo canônico
    # Verificadas contra Magic no Universo Kéos v.0.4.docx (09-03)
    ANCHORS = {
        "§1 Cores — guildas": [
            "Selesnya",
            "Gruul",
            "Rakdos",
            "Dimir",
            "Azorius",
        ],
        "§6 Balizadores": [
            "Foco",
            "Velocidade",
        ],
        "§11 Domínios — Branco": [
            "Alçada da Honra",
            "Alçada da Luz",
        ],
        "§11 Domínios — Verde": [
            "Trilha do Instinto",
            "Trilha da Vegetação",
        ],
        "§11 Domínios — Vermelho": [
            "Desígnio da Ira",
            "Desígnio do Fogo",
        ],
        "§11 Domínios — Preto": [
            "Arte da Dor",
            "Arte da Danação",
        ],
        "§11 Domínios — Azul": [
            "Ramo da Água",
            "Ramo da Contramágica",
        ],
        "§13 Armas": [
            "Montante",
            "Desarmado",
        ],
        "§13 Propriedades Elementais": [
            "Sagrado",
            "Necrotizante",
            "Gélido",
        ],
        "§16 Criaturas": [
            "Toque Mortífero",
            "Amedrontar",
            "Incorpóreo",
        ],
        "§17 Condições": [
            "Congelado",
            "Envenenado",
            "Necrosado",
        ],
        "§18 Combate — movimentação": [
            "Deslocar-se",
            "Esconder-se",
        ],
        "§18 Combate — reações": [
            "Aparar",
            "Esquivar",
        ],
        "§18 Combate — manifestações": [
            "Conjurar",
            "Trucar",
        ],
        "§19 Descanso": [
            "Repousar",
            "Fabricar",
        ],
    }

    drift_count = 0
    results = []

    for section, anchors in ANCHORS.items():
        for anchor in anchors:
            found = check_anchor(all_text, anchor)
            status = "OK" if found else "DRIFT"
            if not found:
                drift_count += 1
            if VERBOSE or not found:
                results.append(f"  [{status}] {section}: '{anchor}'")

    total_anchors = sum(len(v) for v in ANCHORS.values())
    print(f"audit-docx.py — Magic Kéos fidelidade check")
    print(f"DOCX: {DOCX_PATH}")
    print(f"Âncoras verificadas: {total_anchors}")
    print(f"Drift encontrado: {drift_count}")
    print()

    if results:
        for r in results:
            print(r)

    if drift_count > 0:
        print(f"\nFALHA: {drift_count} ancora(s) com drift. Atualize data/regras/ para corresponder ao docx.")
        sys.exit(1)
    else:
        print("OK: Nenhum drift detectado.")
        sys.exit(0)


if __name__ == "__main__":
    main()
